from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfWriter
from zipfile import ZIP_DEFLATED, ZipFile

from genslide_agentscope.content_io import MAX_INPUT_BYTES, MAX_OUTPUT_CHARS, MAX_SLIDE_BODY_CHARS, parse_attachment, render_content


def test_parse_text_attachment_normalizes_preserves_unicode_and_rejects_oversize(tmp_path: Path) -> None:
    source = tmp_path / "notes.TXT"
    source.write_text("  first line 😀  \n\nsecond line", encoding="utf-8")
    result = parse_attachment(source)
    assert result == "first line 😀\n\nsecond line"
    source.write_text("x" * (MAX_OUTPUT_CHARS + 1), encoding="utf-8")
    with pytest.raises(ValueError, match="character limit"):
        parse_attachment(source)


def test_parse_attachment_rejects_oversized_input_file(tmp_path: Path) -> None:
    source = tmp_path / "large.md"
    source.write_bytes(b"x" * (MAX_INPUT_BYTES + 1))
    with pytest.raises(ValueError, match="byte input limit"):
        parse_attachment(source)


def test_parse_pdf_rejects_more_than_200_pages(tmp_path: Path) -> None:
    source = tmp_path / "many-pages.pdf"
    writer = PdfWriter()
    for _ in range(201):
        writer.add_blank_page(width=72, height=72)
    writer.write(source)
    with pytest.raises(ValueError, match="200-page"):
        parse_attachment(source)


def test_parse_docx_rejects_expanded_zip_over_50_mib(tmp_path: Path) -> None:
    source = tmp_path / "expanded.docx"
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (50 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="expanded size"):
        parse_attachment(source)


def test_parse_attachment_rejects_unsupported_type_and_missing_file(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unsupported"):
        parse_attachment(tmp_path / "input.csv")
    with pytest.raises(FileNotFoundError):
        parse_attachment(tmp_path / "missing.txt")


def test_render_presentation_has_cover_section_and_closing(tmp_path: Path) -> None:
    content = {"title": "Plan", "sections": [{"title": "Overview", "body": "Details", "notes": "Speak"}]}
    path = render_content("presentation", content, tmp_path)
    deck = Presentation(path)
    assert path.is_file()
    assert len(deck.slides) == 3
    texts = [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    assert "Plan" in texts and "Overview" in texts and "Thank you" in texts
    body_box = next(shape for shape in deck.slides[1].shapes if shape.has_text_frame and "Details" in shape.text)
    assert body_box.text_frame.word_wrap
    assert body_box.text_frame.paragraphs[0].text == "Details"


def test_render_document_is_editable_when_docx_dependency_is_installed(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    path = render_content("document", {"title": "Plan", "sections": [{"title": "Overview", "body": "Details"}]}, tmp_path)
    assert [paragraph.text for paragraph in Document(path).paragraphs if paragraph.text] == ["Plan", "Overview", "Details"]


def test_document_long_section_is_not_limited_by_slide_capacity(tmp_path: Path) -> None:
    from docx import Document

    body = "长" * (MAX_SLIDE_BODY_CHARS + 1)
    path = render_content("document", {"title": "长文", "sections": [{"title": "章节", "body": body}]}, tmp_path)
    assert Document(path).paragraphs[-1].text == body


def test_render_rejects_invalid_kind_and_content(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="kind"):
        render_content("pdf", {"title": "Plan", "sections": []}, tmp_path)
    with pytest.raises(ValueError, match="sections"):
        render_content("document", {"title": "Plan", "sections": "bad"}, tmp_path)
    with pytest.raises(ValueError, match="at least one"):
        render_content("document", {"title": "Plan", "sections": []}, tmp_path)
    with pytest.raises(ValueError, match="non-empty"):
        render_content("document", {"title": "Plan", "sections": [{"title": "Empty", "body": " "}]}, tmp_path)
    with pytest.raises(ValueError, match="slide limit"):
        render_content("presentation", {"title": "Plan", "sections": [{"title": "Long", "body": "x" * (MAX_SLIDE_BODY_CHARS + 1)}]}, tmp_path)
    with pytest.raises(ValueError, match="limit"):
        render_content("document", {"title": "x" * MAX_OUTPUT_CHARS, "sections": [{"title": "Valid", "body": "body"}]}, tmp_path)


def test_render_document_splits_paragraphs_keeps_lists_and_strips_markdown(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    content = {"title": "Plan", "sections": [{"title": "Overview",
               "body": "First paragraph.\n\nSecond paragraph with **bold** text.\n\n- Item one\n- Item two"}]}
    document = Document(render_content("document", content, tmp_path))
    assert [(p.text, p.style.name) for p in document.paragraphs if p.text] == [
        ("Plan", "Title"),
        ("Overview", "Heading 1"),
        ("First paragraph.", "Normal"),
        ("Second paragraph with bold text.", "Normal"),
        ("Item one", "List Bullet"),
        ("Item two", "List Bullet"),
    ]
    assert any(run.bold for run in document.paragraphs[3].runs)


def test_render_document_excludes_speaker_notes(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    content = {"title": "Plan", "sections": [{"title": "Overview", "body": "Details", "notes": "Speak"}]}
    document = Document(render_content("document", content, tmp_path))
    assert "Speak" not in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_render_presentation_paginates_all_items_and_notes(tmp_path: Path) -> None:
    sections = [{"title": "Items", "body": "\n".join(f"- Item {i}" for i in range(1, 8)), "notes": "Keep"}]
    deck = Presentation(render_content("presentation", {"title": "Plan", "sections": sections}, tmp_path))
    assert len(deck.slides) == 4
    slide_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert all(f"Item {i}" in slide_text for i in range(1, 8))
    assert all("Keep" in slide.notes_slide.notes_text_frame.text for slide in list(deck.slides)[1:-1])


def test_render_presentation_splits_long_items_without_loss(tmp_path: Path) -> None:
    body = "Long item: " + "word " * 900
    deck = Presentation(render_content("presentation", {"title": "Plan", "sections": [{"title": "Long", "body": body}]}, tmp_path))
    rendered = "".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert rendered.count("word") == 900
    body_boxes = [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame and "word" in shape.text]
    assert max(map(len, body_boxes)) <= 60


def test_render_document_handles_mixed_prose_and_lists(tmp_path: Path) -> None:
    from docx import Document

    body = "Intro prose\n- First\n- Second\nClosing prose"
    document = Document(render_content("document", {"title": "Plan", "sections": [{"title": "Mixed", "body": body}]}, tmp_path))
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.text] == ["Plan", "Mixed", "Intro prose", "First", "Second", "Closing prose"]


def test_render_document_tables_are_editable(tmp_path: Path) -> None:
    from docx import Document

    body = "| Name | Value |\n| --- | --- |\n| A | 1 |"
    document = Document(render_content("document", {"title": "Plan", "sections": [{"title": "Table", "body": body}]}, tmp_path))
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [["Name", "Value"], ["A", "1"]]


def test_parse_docx_preserves_paragraph_table_order_and_cells(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"
    document.add_paragraph("After")
    document.save(source)
    assert parse_attachment(source) == "Before\n\nLeft | Right\n\nAfter"
