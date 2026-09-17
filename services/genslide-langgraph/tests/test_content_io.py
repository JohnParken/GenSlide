from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfWriter
from zipfile import ZIP_DEFLATED, ZipFile

from genslide_langgraph.content_io import MAX_INPUT_BYTES, MAX_OUTPUT_CHARS, MAX_SLIDE_BODY_CHARS, parse_attachment, render_content


def test_parse_text_attachment_normalizes_preserves_unicode_and_rejects_oversize(tmp_path: Path) -> None:
    source = tmp_path / "notes.TXT"
    source.write_text("  first line 😀  \n\nsecond line", encoding="utf-8")
    result = parse_attachment(source)
    assert result == "first line 😀\n\nsecond line"
    source.write_text("x" * (MAX_OUTPUT_CHARS + 1), encoding="utf-8")
    with pytest.raises(ValueError, match="character limit"):
        parse_attachment(source)


def test_parse_attachment_rejects_oversized_input_file(tmp_path: Path) -> None:
    source = tmp_path / "large.md"
    source.write_bytes(b"x" * (MAX_INPUT_BYTES + 1))
    with pytest.raises(ValueError, match="byte input limit"):
        parse_attachment(source)


def test_parse_pdf_rejects_more_than_200_pages(tmp_path: Path) -> None:
    source = tmp_path / "many-pages.pdf"
    writer = PdfWriter()
    for _ in range(201):
        writer.add_blank_page(width=72, height=72)
    writer.write(source)
    with pytest.raises(ValueError, match="200-page"):
        parse_attachment(source)


def test_parse_docx_rejects_expanded_zip_over_50_mib(tmp_path: Path) -> None:
    source = tmp_path / "expanded.docx"
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (50 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="expanded size"):
        parse_attachment(source)


def test_parse_attachment_rejects_unsupported_type_and_missing_file(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unsupported"):
        parse_attachment(tmp_path / "input.csv")
    with pytest.raises(FileNotFoundError):
        parse_attachment(tmp_path / "missing.txt")


def test_render_presentation_has_cover_section_and_closing(tmp_path: Path) -> None:
    content = {"title": "Plan", "sections": [{"title": "Overview", "body": "Details", "notes": "Speak"}]}
    path = render_content("presentation", content, tmp_path)
    deck = Presentation(path)
    assert path.is_file()
    assert len(deck.slides) == 3
    texts = [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    assert "Plan" in texts and "Overview" in texts and "Thank you" in texts
    body_box = next(shape for shape in deck.slides[1].shapes if shape.has_text_frame and "Details" in shape.text)
    assert body_box.text_frame.word_wrap
    assert body_box.text_frame.paragraphs[0].text == "Details"


def test_render_document_is_editable_when_docx_dependency_is_installed(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    path = render_content("document", {"title": "Plan", "sections": [{"title": "Overview", "body": "Details"}]}, tmp_path)
    assert [paragraph.text for paragraph in Document(path).paragraphs if paragraph.text] == ["Plan", "Overview", "Details"]


def test_document_long_section_is_not_limited_by_slide_capacity(tmp_path: Path) -> None:
    from docx import Document

    body = "长" * (MAX_SLIDE_BODY_CHARS + 1)
    path = render_content("document", {"title": "长文", "sections": [{"title": "章节", "body": body}]}, tmp_path)
    assert Document(path).paragraphs[-1].text == body


def test_render_rejects_invalid_kind_and_content(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="kind"):
        render_content("pdf", {"title": "Plan", "sections": []}, tmp_path)
    with pytest.raises(ValueError, match="sections"):
        render_content("document", {"title": "Plan", "sections": "bad"}, tmp_path)
    with pytest.raises(ValueError, match="at least one"):
        render_content("document", {"title": "Plan", "sections": []}, tmp_path)
    with pytest.raises(ValueError, match="non-empty"):
        render_content("document", {"title": "Plan", "sections": [{"title": "Empty", "body": " "}]}, tmp_path)
    with pytest.raises(ValueError, match="slide limit"):
        render_content("presentation", {"title": "Plan", "sections": [{"title": "Long", "body": "x" * (MAX_SLIDE_BODY_CHARS + 1)}]}, tmp_path)
    with pytest.raises(ValueError, match="limit"):
        render_content("document", {"title": "x" * MAX_OUTPUT_CHARS, "sections": [{"title": "Valid", "body": "body"}]}, tmp_path)
