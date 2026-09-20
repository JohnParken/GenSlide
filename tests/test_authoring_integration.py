"""Boundary regressions for shared writing budgets and complete attachments."""
import asyncio
import io
import json
from unittest.mock import AsyncMock

from docx import Document

from frontend.autonomous_agent import AutonomousAgent, detect_target_kind
from frontend.chat_context import read_attachment_bytes, updated_requirements, RequirementUpdate
from genslide_agentscope.authoring import explicit_length, length_bounds, length_target


def test_latest_affirmative_length_and_limits_are_preserved():
    assert explicit_length("不要5000字，改成3000字") == "3000字"
    assert explicit_length("取消3000字的限制") is None
    assert explicit_length("写一份不超过3000字的报告") == "不超过3000字"
    assert length_bounds({"length": "不超过3000字"}) == (0, 3000)
    assert length_target({"length": "3-5千字"}) == 4000
    assert length_target({"length": "1w字"}) == 10000


def test_new_constraint_does_not_erase_previous_constraints():
    memory = {"requirements": {"constraints": "不要编造数据"}}
    added = updated_requirements(memory, {"constraints": RequirementUpdate(value="不要术语", evidence="不要术语")}, "另外不要术语")
    assert added["constraints"] == "不要编造数据\n不要术语"
    removed = updated_requirements({"requirements": added}, {
        "constraints": RequirementUpdate(value=None, evidence="取消不要术语")
    }, "取消不要术语")
    assert removed["constraints"] == "不要编造数据"


def test_attachment_reader_keeps_gbk_and_docx_table_order():
    assert read_attachment_bytes("材料.TXT", "预算为73万元".encode("gb18030")) == "预算为73万元"
    doc = Document()
    doc.add_paragraph("表格之前")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "负责人"
    table.cell(0, 1).text = "林岚"
    doc.add_paragraph("表格之后")
    buffer = io.BytesIO()
    doc.save(buffer)
    text = read_attachment_bytes("材料.DOCX", buffer.getvalue())
    assert text.index("表格之前") < text.index("负责人 | 林岚") < text.index("表格之后")


def test_ppt_followup_inherits_actual_format_without_switching_on_material_words():
    assert detect_target_kind("结合材料补充第二页", default="presentation") == "presentation"
    assert detect_target_kind("做一份季度汇报PPT") == "presentation"
    assert detect_target_kind("把PPT转成Word") == "document"
    agent = AutonomousAgent()
    agent._call_llm = AsyncMock(return_value='{"skill":"reply","reply_text":"已了解"}')
    result = asyncio.run(agent.astep("第二页还有什么建议", memory={"target_kind": "presentation"}))
    assert result.target_kind == "presentation"


def test_confirmed_requirement_survives_ten_turns_without_summary_loss():
    history = []
    for index in range(10):
        history.extend([{"role": "user", "text": "面向董事会" if index == 0 else f"讨论{index}"},
                        {"role": "assistant", "text": "继续讨论"}])
    agent = AutonomousAgent()
    agent._call_llm = AsyncMock(return_value='{"skill":"reply","reply_text":"面向董事会"}')
    memory = {"requirements": {"audience": "董事会", "constraints": "不要编造数字"}}
    result = asyncio.run(agent.astep("回顾要求", history=history, memory=memory))
    context = json.loads(agent._call_llm.call_args.args[1])
    assert "面向董事会" in context["conversation"]
    assert result.memory["requirements"] == memory["requirements"]


def test_failed_render_does_not_mutate_input_draft_or_memory():
    from unittest.mock import patch

    original = {"title": "标题", "sections": [{"title": "章节", "body": "旧文", "notes": ""}]}
    memory = {"requirements": {"style": "务实"}}
    agent = AutonomousAgent()
    agent._call_llm = AsyncMock(side_effect=[
        '{"skill":"revise","section_ids":["sec-1"]}',
        '{"sections":[{"title":"章节","body":"新文","notes":""}]}',
    ])
    with patch("frontend.autonomous_agent.render_content", side_effect=ValueError("render failed")):
        try:
            asyncio.run(agent.astep("修改第一章", current_content=original, memory=memory))
        except ValueError as exc:
            assert str(exc) == "render failed"
        else:
            raise AssertionError("render failure must be explicit")
    assert original["sections"][0]["body"] == "旧文"
    assert memory == {"requirements": {"style": "务实"}}


def test_new_outline_is_used_for_generation_even_when_an_old_draft_exists():
    agent = AutonomousAgent()
    draft = {"title": "旧标题", "sections": [{"title": "旧章节", "body": "旧文"}]}
    outline = {"title": "新标题", "nodes": [{"node_id": "new", "title": "新章节"}]}
    agent._call_llm = AsyncMock(side_effect=[
        '{"skill":"generate","reply_text":"开始"}',
        '{"title":"新标题","sections":[{"title":"新章节","body":"新的完整正文","notes":""}]}',
    ])
    result = asyncio.run(agent.astep("按新大纲生成纯文本", current_content=draft, current_outline=outline))
    assert result.content["sections"][0]["title"] == "新章节"
    request = json.loads(agent._call_llm.call_args_list[1].args[1])
    assert request["outline"] == outline
