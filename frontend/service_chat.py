"""GenSlide AI Writing Studio - Apple-inspired Workbench with Expert Advisory Council.

Features:
1. Creative Exploration Mode with persistent context and optional focused questions;
2. Professional Workflow Mode (专业模式) with model-driven intelligent state transitions;
3. WorkBuddy-style Expert Advisory Council (专家团) with vivid personas and specialized guidance;
4. Apple-grade frosted glass aesthetics and live artifact inspector.
"""
from __future__ import annotations

import io
import hashlib
import json
import os
import uuid
from copy import deepcopy
from pathlib import Path
import uuid
from zipfile import BadZipFile, ZipFile
import streamlit as st

try:
    from frontend.service_chat_client import ChatClient, ChatClientError, ChatState
    from frontend.autonomous_agent import AutonomousAgent, detect_target_kind, route_professional_intent
    from frontend.expert_council import get_expert, list_all_experts, ExpertProfile
    from frontend.chat_context import read_attachment_bytes
    from frontend.chat_session import load_session, save_session
except ImportError:  # streamlit run frontend/service_chat.py
    from service_chat_client import ChatClient, ChatClientError, ChatState
    from autonomous_agent import AutonomousAgent, detect_target_kind, route_professional_intent
    from expert_council import get_expert, list_all_experts, ExpertProfile
    from chat_context import read_attachment_bytes
    from chat_session import load_session, save_session

def rerun_workbench():
    save_session(st.session_state)
    st.rerun()


st.set_page_config(
    page_title="GenSlide · AI 写作工作台",
    page_icon="✨",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# 附件预检校验辅助函数
# ---------------------------------------------------------------------------
def validate_attachment(filename: str, content: bytes) -> tuple[bool, str]:
    """Pre-flight check for attachment validity before submitting to BFF/backend."""
    if not content or len(content) == 0:
        return False, "所选文件内容为空（0 字节），无法提取有效文本。"
    if len(content) > 20 * 1024 * 1024:
        return False, "文件大小超过 20MB 上限。"

    suffix = Path(filename).suffix.lower()
    if suffix not in {".txt", ".md", ".pdf", ".docx"}:
        return False, f"不支持的文件后缀 `{suffix}`，仅支持 .txt, .md, .pdf, .docx。"

    if suffix in {".txt", ".md"}:
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content.decode("gb18030")
            except UnicodeDecodeError:
                return False, "文本文件编码不支持，请确保为 UTF-8 或 GBK/GB18030 编码。"
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > 200:
                return False, f"PDF 页数 ({len(reader.pages)}) 超过 200 页限制。"
            has_text = any((page.extract_text() or "").strip() for page in reader.pages)
            if not has_text:
                return False, (
                    "该 PDF 文件未检测到文本图层（疑似纯图片或扫描件）。\n"
                    "后端需要可直接提取的文字（当前未集成 OCR 图文识别）。请上传包含可复制文字的 PDF，或使用 TXT/Word 文档。"
                )
        except Exception as exc:
            return False, f"PDF 解析失败或损坏: {exc}"
    elif suffix == ".docx":
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if any(entry.flag_bits & 0x1 for entry in archive.infolist()):
                    return False, "不支持受密码加密保护的 DOCX 文件。"
        except BadZipFile:
            return False, "无效的 DOCX 文件（若是旧版 .doc 格式，请在 Word 中另存为 .docx 后再上传）。"
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            has_text = any(p.text.strip() for p in doc.paragraphs) or any(
                cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells
            )
            if not has_text:
                return False, "DOCX 文档内未找到任何有效正文段落或表格文字。"
        except Exception as exc:
            return False, f"DOCX 解析失败: {exc}"

    return True, ""


# ---------------------------------------------------------------------------
# 网关与客户端初始化（默认无感开箱即用，支持通过耳机🎧控制台随时热调整）
# ---------------------------------------------------------------------------
default_bff = os.getenv("GENSLIDE_BFF_URL", "http://localhost:8010")
default_token = os.getenv("GENSLIDE_SERVICE_TOKEN", "local-development-token-at-least-32-characters")
default_engine = os.getenv("GENSLIDE_ENGINE", "agentscope")
default_service = os.getenv("GENSLIDE_SERVICE_URL", "http://localhost:8002")

if "bff_url" not in st.session_state:
    st.session_state.bff_url = default_bff
if "bff_token" not in st.session_state:
    st.session_state.bff_token = default_token
if "service_url" not in st.session_state:
    st.session_state.service_url = default_service
if "engine" not in st.session_state:
    st.session_state.engine = default_engine

client = ChatClient(
    st.session_state.bff_url,
    st.session_state.service_url,
    st.session_state.bff_token,
    engine=st.session_state.engine,
)
bff = st.session_state.bff_url
token = st.session_state.bff_token
service = st.session_state.service_url
engine = st.session_state.engine

# ---------------------------------------------------------------------------
# 全局多会话管理 (Multi-Session Storage in session_state)
# ---------------------------------------------------------------------------
if "sessions_store" not in st.session_state:
    # 默认新建首个会话
    first_sid = uuid.uuid4().hex[:8]
    st.session_state.sessions_store = {
        first_sid: {
            "title": "新创作对话",
            "created_at": "刚刚",
            "chat_state": ChatState(engine=engine),
            "messages": [],
            "downloads": [],
            "current_attachment": None,
            "active_follow_ups": [],
            "uploader_key": 0,
        }
    }
    st.session_state.current_session_id = first_sid

if "current_session_id" not in st.session_state or st.session_state.current_session_id not in st.session_state.sessions_store:
    st.session_state.current_session_id = next(iter(st.session_state.sessions_store.keys()))

# 写作工作台抽屉/分栏展开开关（默认展开工作台，点击可切换三栏或收拢）
if "show_workbench" not in st.session_state:
    st.session_state.show_workbench = True

# ---------------------------------------------------------------------------
# 侧边栏：用户会话栏 (User Sessions Sidebar)
# ---------------------------------------------------------------------------
with st.sidebar:
    col_new_btn, col_refresh_btn = st.columns([4, 1])
    with col_new_btn:
        if st.button("➕ 新建会话", use_container_width=True, type="primary"):
            new_sid = uuid.uuid4().hex[:8]
            st.session_state.sessions_store[new_sid] = {
                "title": f"新会话 {len(st.session_state.sessions_store) + 1}",
                "created_at": "刚刚",
                "chat_state": ChatState(engine=engine),
                "messages": [],
                "downloads": [],
                "current_attachment": None,
                "active_follow_ups": [],
                "uploader_key": 0,
            }
            st.session_state.current_session_id = new_sid
            rerun_workbench()

    with col_refresh_btn:
        if st.button("🔄", help="刷新技能与专家团"):
            try:
                client.reload_skills()
            except Exception:
                pass
            st.toast("专家团已刷新")
            rerun_workbench()

    st.markdown("#### 💬 历史创作会话")

    # 遍历渲染会话列表
    all_sids = list(st.session_state.sessions_store.keys())
    for sid in reversed(all_sids):
        s_data = st.session_state.sessions_store[sid]
        s_title = s_data.get("title", f"会话 {sid}")
        # 如果当前大纲或正文有标题，动态同步
        d_title = (s_data["chat_state"].content or {}).get("title") or (s_data["chat_state"].draft or {}).get("title")
        if d_title and s_title.startswith("新会话"):
            s_title = d_title[:14] + ("..." if len(d_title) > 14 else "")
            s_data["title"] = s_title

        is_active = sid == st.session_state.current_session_id
        display_label = f"📌 {s_title}" if is_active else f"💭 {s_title}"

        col_s_item, col_s_del = st.columns([5, 1])
        with col_s_item:
            if st.button(
                display_label,
                key=f"sess_btn_{sid}",
                use_container_width=True,
                type="primary" if is_active else "secondary",
            ):
                st.session_state.current_session_id = sid
                rerun_workbench()
        with col_s_del:
            if len(st.session_state.sessions_store) > 1:
                if st.button("✕", key=f"del_sess_{sid}", help="删除此会话"):
                    del st.session_state.sessions_store[sid]
                    if st.session_state.current_session_id == sid:
                        st.session_state.current_session_id = next(iter(st.session_state.sessions_store.keys()))
                    rerun_workbench()

    st.divider()

    # 坐镇专家与交付形态收纳配置
    with st.expander("✍️ 创作顾问与参数配置", expanded=False):
        # 1. 坐镇专家顾问 (当前创作角色)
        discovered_skills = []
        try:
            discovered_skills = client.list_skills()
        except Exception:
            discovered_skills = []

        if not discovered_skills:
            try:
                if "autonomous_agent" not in st.session_state:
                    st.session_state.autonomous_agent = AutonomousAgent()
                discovered_skills = st.session_state.autonomous_agent.list_skills()
            except Exception:
                discovered_skills = []

        all_experts = list_all_experts(discovered_skills)
        expert_options = [f"{exp.avatar} {exp.name}" for exp in all_experts]
        expert_map = {opt: exp for opt, exp in zip(expert_options, all_experts)}

        selected_expert_label = st.selectbox(
            "坐镇顾问 (Expert)",
            expert_options,
            key=f"expert_{st.session_state.current_session_id}",
            index=next((i for i, expert in enumerate(all_experts)
                        if expert.id == st.session_state.sessions_store[st.session_state.current_session_id].get("expert_id", "default")), 0),
            help="选派各细分领域专家，为您提供针对性行文规范与建议",
        )
        current_expert = expert_map[selected_expert_label]
        skill_id = current_expert.id if current_expert.id != "default" else ""

        # 2. 目标交付形态
        target = st.selectbox(
            "交付形态 (Target)",
            ["document", "presentation", "writing"],
            key=f"target_{st.session_state.current_session_id}",
            index=["document", "presentation", "writing"].index(
                st.session_state.sessions_store[st.session_state.current_session_id].get("ui_target", "document")),
            format_func=lambda x: {
                "document": "📄 深度长文 / Word (.docx)",
                "presentation": "📊 方案演示 / PPT (.pptx)",
                "writing": "📝 纯文本 / 草稿",
            }.get(x, x),
        )

        # 3. 创作协同模式
        interaction_mode = st.radio(
            "创作模式 (Mode)",
            ["🌟 自由灵感模式 (自主规划/选项直选)", "📐 结构工作流 (澄清->大纲->成稿)"],
            index=0,
        )
        is_autonomous = interaction_mode.startswith("🌟")

        # 4. 视觉主题
        ui_theme = st.selectbox(
            "画布主题 (Theme)",
            ["☀️ 经典浅色 (Doubao Paper White)", "🌙 暗黑深邃 (Dark Canvas)"],
            index=0,
        )
        is_dark = ui_theme.startswith("🌙")


# ---------------------------------------------------------------------------
# Apple 视觉主题色彩动态计算 (Theme Palette)
# ---------------------------------------------------------------------------
if is_dark:
    app_bg = "#000000"
    text_main = "#f5f5f7"
    text_sub = "#86868b"
    text_caption = "#a1a1a6"
    sidebar_bg = "#161618"
    sidebar_border = "rgba(255, 255, 255, 0.1)"
    card_bg = "rgba(28, 28, 30, 0.65)"
    card_border = "rgba(255, 255, 255, 0.08)"
    card_shadow = "0 8px 30px rgba(0, 0, 0, 0.4)"
    header_bg = "rgba(28, 28, 30, 0.75)"
    header_border = "rgba(255, 255, 255, 0.1)"
    header_title_color = "#ffffff"
    tag_bg = "rgba(255, 255, 255, 0.08)"
    tag_border = "rgba(255, 255, 255, 0.1)"
    tag_text = "#d1d1d6"
    chat_bubble_bg = "rgba(28, 28, 30, 0.6)"
    chat_bubble_border = "rgba(255, 255, 255, 0.08)"
    btn_bg = "rgba(255, 255, 255, 0.06)"
    btn_border = "rgba(255, 255, 255, 0.12)"
    btn_text = "#ffffff"
    banner_bg = "rgba(28, 28, 30, 0.6)"
    banner_border = "rgba(255, 255, 255, 0.08)"
else:
    app_bg = "#fbfbfd"
    text_main = "#1d1d1f"
    text_sub = "#86868b"
    text_caption = "#424245"
    sidebar_bg = "#f5f5f7"
    sidebar_border = "rgba(0, 0, 0, 0.08)"
    card_bg = "#ffffff"
    card_border = "rgba(0, 0, 0, 0.08)"
    card_shadow = "0 2px 10px rgba(0, 0, 0, 0.04), 0 1px 3px rgba(0, 0, 0, 0.02)"
    header_bg = "rgba(255, 255, 255, 0.85)"
    header_border = "rgba(0, 0, 0, 0.08)"
    header_title_color = "#1d1d1f"
    tag_bg = "#f5f5f7"
    tag_border = "rgba(0, 0, 0, 0.06)"
    tag_text = "#1d1d1f"
    chat_bubble_bg = "#ffffff"
    chat_bubble_border = "rgba(0, 0, 0, 0.08)"
    btn_bg = "#ffffff"
    btn_border = "rgba(0, 0, 0, 0.12)"
    btn_text = "#1d1d1f"
    banner_bg = "rgba(0, 113, 227, 0.04)"
    banner_border = "rgba(0, 113, 227, 0.15)"


# ---------------------------------------------------------------------------
# Apple Design System CSS 注入
# ---------------------------------------------------------------------------
APPLE_CSS = f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

/* Global Apple typography and base styling */
html, body, [data-testid="stAppViewContainer"], .main {{
    font-family: -apple-system, BlinkMacSystemFont, "SF Pro Display", "SF Pro Text", "Inter", "Helvetica Neue", sans-serif !important;
    color: {text_main} !important;
    background-color: {app_bg} !important;
}}

[data-testid="stSidebar"] {{
    background-color: {sidebar_bg} !important;
    border-right: 1px solid {sidebar_border} !important;
}}

/* Header bar */
.apple-header {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 14px 20px;
    background: {header_bg};
    backdrop-filter: blur(25px);
    -webkit-backdrop-filter: blur(25px);
    border: 1px solid {header_border};
    border-radius: 18px;
    margin-bottom: 20px;
    box-shadow: {card_shadow};
}}

.apple-title-wrap {{
    display: flex;
    align-items: center;
    gap: 12px;
}}

.apple-logo-badge {{
    font-size: 24px;
    background: rgba(0, 113, 227, 0.08);
    border-radius: 12px;
    width: 44px;
    height: 44px;
    display: flex;
    align-items: center;
    justify-content: center;
    border: 1px solid rgba(0, 113, 227, 0.18);
}}

.apple-main-title {{
    font-size: 20px;
    font-weight: 700;
    letter-spacing: -0.5px;
    color: {header_title_color};
    margin: 0;
}}

.apple-sub-title {{
    font-size: 12px;
    color: {text_sub};
    margin-top: 2px;
}}

.status-dot {{
    width: 8px;
    height: 8px;
    border-radius: 50%;
    display: inline-block;
    background-color: #34c759;
    box-shadow: 0 0 8px #34c759;
    margin-right: 6px;
}}

.apple-pill {{
    display: inline-flex;
    align-items: center;
    padding: 4px 12px;
    border-radius: 9999px;
    font-size: 12px;
    font-weight: 500;
    background: rgba(0, 113, 227, 0.08);
    color: #0071e3;
    border: 1px solid rgba(0, 113, 227, 0.2);
}}

/* Streamlit button refinement */
div.stButton > button {{
    border-radius: 12px !important;
    font-weight: 500 !important;
    transition: all 0.2s cubic-bezier(0.16, 1, 0.3, 1) !important;
    border: 1px solid {btn_border} !important;
    background: {btn_bg} !important;
    color: {btn_text} !important;
}}

div.stButton > button:hover {{
    border-color: #0071e3 !important;
    color: #0071e3 !important;
    box-shadow: 0 4px 16px rgba(0, 113, 227, 0.15) !important;
    transform: translateY(-1px) !important;
}}

/* Primary buttons with Apple Blue gradient */
div.stButton > button[kind="primary"] {{
    background: linear-gradient(135deg, #0071e3 0%, #0077ed 100%) !important;
    border: none !important;
    color: #ffffff !important;
}}

div.stButton > button[kind="primary"]:hover {{
    background: linear-gradient(135deg, #0077ed 0%, #2997ff 100%) !important;
}}

/* Chat message bubbles */
[data-testid="stChatMessage"] {{
    background: {chat_bubble_bg} !important;
    border: 1px solid {chat_bubble_border} !important;
    border-radius: 16px !important;
    padding: 14px 18px !important;
    margin-bottom: 12px !important;
    box-shadow: 0 2px 8px rgba(0, 0, 0, 0.02) !important;
}}

/* Tabs styling */
[data-testid="stTabs"] button {{
    border-radius: 10px !important;
    font-weight: 600 !important;
    color: {text_sub} !important;
}}

[data-testid="stTabs"] button[aria-selected="true"] {{
    color: #0071e3 !important;
    border-bottom-color: #0071e3 !important;
}}

/* 豆包式写文章右侧白板画布 (Doubao Paper Canvas) */
.doubao-paper-canvas {{
    background: {card_bg};
    border: 1px solid {card_border};
    border-radius: 18px;
    padding: 32px 36px;
    box-shadow: 0 4px 24px rgba(0, 0, 0, 0.04), 0 1px 4px rgba(0, 0, 0, 0.02);
    min-height: 520px;
    margin-bottom: 20px;
}}

.doubao-doc-header {{
    border-bottom: 1px solid {card_border};
    padding-bottom: 18px;
    margin-bottom: 22px;
}}

.doubao-doc-title {{
    font-size: 24px;
    font-weight: 700;
    color: {text_main};
    letter-spacing: -0.5px;
    line-height: 1.3;
}}

.doubao-doc-meta {{
    display: flex;
    align-items: center;
    gap: 12px;
    font-size: 12px;
    color: {text_sub};
    margin-top: 8px;
}}

.doubao-section-block {{
    background: rgba(0, 113, 227, 0.02);
    border-left: 3px solid #0071e3;
    padding: 14px 18px;
    border-radius: 0 12px 12px 0;
    margin-bottom: 18px;
}}

.doubao-section-title {{
    font-size: 16px;
    font-weight: 600;
    color: {text_main};
    margin-bottom: 8px;
}}

.doubao-section-body {{
    font-size: 14.5px;
    line-height: 1.7;
    color: {text_caption};
}}
</style>
"""
st.markdown(APPLE_CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# 从多会话存储池同步当前活跃会话状态 (Sync Active Session from store)
# ---------------------------------------------------------------------------
active_sess = load_session(st.session_state)
state: ChatState = active_sess["chat_state"]
if active_sess.get("ui_target", target) != target:
    state.autonomous_memory["target_kind"] = target
active_sess["ui_target"] = target
active_sess["expert_id"] = skill_id or "default"
# A follow-up inherits the actual last output format until the user changes it.
conversation_target = state.autonomous_memory.get("target_kind") or (state.content or {}).get("target_kind") or (state.draft or {}).get("target_kind") or target

if "autonomous_agent" not in st.session_state:
    st.session_state.autonomous_agent = AutonomousAgent()

# 侧边栏底部简要展示当前会话 ID
with st.sidebar:
    st.caption(f"🆔 当前会话: `{state.session[:8]}` | 版本: `v{state.session_version}`")


# ---------------------------------------------------------------------------
# 阶段解析函数
# ---------------------------------------------------------------------------
def compute_stage(state: ChatState) -> tuple[int, str]:
    if state.content:
        return 4, "已生成交付物 (Completed)"
    if state.guidance.get("stage") == "confirmed" or (state.draft and state.draft.get("confirmed_hash")):
        return 3, "大纲已锁定确认 (Confirmed)"
    if state.draft:
        return 2, f"大纲制定中 v{state.draft.get('outline_version', 1)} (Outline)"
    return 1, "需求澄清中 (Clarify)"


current_stage_idx, current_stage_name = compute_stage(state)


# ---------------------------------------------------------------------------
# 格式化助手回复
# ---------------------------------------------------------------------------
def format_assistant_message(state: ChatState, operation: str) -> str:
    lines = []

    # 1. 结构化大纲生成或修订
    if operation in {"create_outline", "revise_outline"} and state.draft:
        draft = state.draft
        lines.append(f"📑 **大纲已成功生成：【{draft.get('title', '未命名大纲')}】** (版本: `v{draft.get('outline_version', 1)}`)")
        lines.append("\n**章节结构清单：**")
        nodes = draft.get("nodes", [])
        for idx, node in enumerate(nodes, 1):
            lines.append(f"{idx}. 📌 **{node.get('title')}** `(id: {node.get('node_id')})`")
        lines.append(
            "\n👉 **智能流转建议**：\n"
            "- 若大纲结构满意，直接在下方输入「确认大纲」或「就按这个大纲出正文」，AI 将自动为您锁定并生成；\n"
            "- 若需调整章节，直接输入修改意见（例如：*增加一节案例分析* 或 *第三节改成落地实践*）。"
        )
        return "\n".join(lines)

    # 2. 确认大纲
    if operation == "confirm_outline":
        conf_hash = (state.draft or {}).get("confirmed_hash") or ""
        short_hash = f" `({conf_hash[:8]})`" if conf_hash else ""
        ver = (state.draft or {}).get("outline_version", 1)
        return (
            f"🔒 **大纲已正式锁定确认！**{short_hash}\n\n"
            f"大纲结构现已刚性锁定（版本 `v{ver}`）。\n"
            f"直接输入「开始生成」或点击下方按钮，即可启动交付物渲染！"
        )

    # 3. 交付生成
    if operation == "generate" and state.content:
        content = state.content
        sections = content.get("sections", [])
        lines.append("🎉 **完整内容交付成功！**")
        lines.append(f"- **主标题**：【{content.get('title', '无标题')}】")
        lines.append(f"- **章节总数**：共 {len(sections)} 个完整章节")
        if content.get("summary"):
            lines.append(f"- **内容摘要**：{content.get('summary')}")
        lines.append("\n👉 您可以在右侧面板查看 **「📖 交付正文」**，或前往 **「💾 文件下载」** 获取导出的文件。")
        return "\n\n".join(lines)

    # 4. 大纲解释 (explain_outline)
    if operation == "explain_outline":
        ans = state.answer.strip() if state.answer else "暂无针对该大纲的解释说明。"
        return f"💬 **大纲解析说明**：\n\n{ans}"

    # 5. 澄清与讨论阶段 (clarify)
    if state.answer and state.answer.strip():
        lines.append(state.answer.strip())

    guidance = state.guidance or {}

    # 阶段指引
    summary = guidance.get("summary", "").strip()
    if summary and summary not in (state.answer or ""):
        lines.append(f"**📋 阶段指引**\n\n{summary}")

    # 待明确问题
    questions = guidance.get("questions", [])
    if questions:
        q_lines = ["\n**❓ 待明确问题：**"]
        for idx, q in enumerate(questions, 1):
            q_lines.append(f"{idx}. {q.get('text', '')}")
            opts = q.get("options", [])
            if opts:
                opts_str = " | ".join(f"`{opt}`" for opt in opts)
                q_lines.append(f"   *可选方案*: {opts_str}")
        lines.append("\n".join(q_lines))

    # 建议提案
    proposals = guidance.get("proposals", [])
    if proposals:
        p_lines = ["\n**💡 助手建议：**"]
        for idx, p in enumerate(proposals, 1):
            reason = f" *（{p['reason']}）*" if p.get("reason") else ""
            p_lines.append(f"{idx}. **{p.get('field', '建议')}**: `{p.get('value', '')}`{reason}")
        lines.append("\n".join(p_lines))

    if not lines:
        return "✅ 操作已成功完成，请继续输入下一步指令。"

    return "\n\n".join(lines)


# ---------------------------------------------------------------------------
# 豆包风格顶栏：沉浸式文档导航与 🎧 高级设置收纳
# ---------------------------------------------------------------------------
doc_title = (state.content or {}).get("title") or (state.draft or {}).get("title") or "未命名创作工程"
doc_sections = (state.content or {}).get("sections", [])
total_words = sum(len(s.get("body", "")) for s in doc_sections)
if not total_words and state.content and state.content.get("summary"):
    total_words = len(state.content.get("summary"))

mode_label = "🌟 自由灵感" if is_autonomous else "📐 结构工作流"
mode_badge_style = (
    "background: rgba(255, 255, 255, 0.08); color: #e5e5ea; border: 1px solid rgba(255, 255, 255, 0.1);"
    if is_dark
    else "background: #f2f2f7; color: #1d1d1f; border: 1px solid rgba(0, 0, 0, 0.08);"
)

# 顶部导航行：左侧品牌与文档名，右侧坐镇专家与 🎧 高级设置
col_nav_left, col_nav_right = st.columns([7, 3])
with col_nav_left:
    st.markdown(
        f"""
        <div style="display: flex; align-items: center; gap: 14px; padding: 4px 0 8px 0;">
            <div class="apple-logo-badge" style="width: 38px; height: 38px; font-size: 20px; border-radius: 10px;">✨</div>
            <div>
                <div style="font-size: 17px; font-weight: 700; color: {text_main}; letter-spacing: -0.3px; display: flex; align-items: center; gap: 8px;">
                    {doc_title}
                    <span style="font-size: 11px; font-weight: 500; color: #0071e3; background: rgba(0, 113, 227, 0.08); padding: 2px 8px; border-radius: 9999px; border: 1px solid rgba(0, 113, 227, 0.2);">
                        {current_stage_name}
                    </span>
                </div>
                <div style="font-size: 11px; color: {text_sub}; margin-top: 2px; display: flex; align-items: center; gap: 12px;">
                    <span>字数统计: <b>{total_words}</b> 字</span>
                    <span>•</span>
                    <span>目标形态: <b>{target}</b></span>
                    <span>•</span>
                    <span>坐镇专家: <b>{current_expert.avatar} {current_expert.name}</b></span>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with col_nav_right:
    col_status, col_headphone = st.columns([1, 1])
    with col_status:
        st.markdown(
            f"""
            <div style="text-align: right; padding-top: 10px;">
                <span style="display: inline-flex; align-items: center; padding: 4px 10px; border-radius: 9999px; font-size: 11px; font-weight: 500; {mode_badge_style}">
                    <span class="status-dot"></span> {mode_label}
                </span>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with col_headphone:
        # 🎧 专用高级配置抽屉（默认完全折叠对普通用户隐蔽，仅供调试或网关配置使用）
        with st.popover("🎧 高级控制", help="服务网关配置与系统级底层调试控制台 (默认已隐藏)"):
            st.markdown("#### 🎧 服务网关与通信配置")
            st.caption("以下技术参数已自动调优，常规写作无需修改：")
            new_bff = st.text_input("BFF URL", value=st.session_state.bff_url)
            new_token = st.text_input("Token", value=st.session_state.bff_token, type="password")
            new_service = st.text_input("Agent URL", value=st.session_state.service_url)
            new_engine = st.selectbox("Engine", ["agentscope", "mock"], index=0 if st.session_state.engine == "agentscope" else 1)

            if st.button("💾 保存并应用网关连接", use_container_width=True, type="primary"):
                st.session_state.bff_url = new_bff
                st.session_state.bff_token = new_token
                st.session_state.service_url = new_service
                st.session_state.engine = new_engine
                st.toast("✅ 网关通信参数已成功更新！")
                rerun_workbench()

            st.divider()
            with st.expander("🛠️ 查看系统底层状态数据", expanded=False):
                st.caption("Guidance:")
                st.json(state.guidance or {})
                st.caption("Draft:")
                st.json(state.draft or {})
                st.caption("Content:")
                st.json(state.content or {})

if not is_autonomous:
    # 顶部阶段 Stepper 状态条
    col_s1, col_s2, col_s3, col_s4 = st.columns(4)
    with col_s1:
        if current_stage_idx == 1:
            st.info("🔹 **1. 需求澄清**")
        elif current_stage_idx > 1:
            st.success("✔ **1. 需求澄清**")
        else:
            st.caption("1. 需求澄清")
    with col_s2:
        if current_stage_idx == 2:
            st.info("🔹 **2. 制定大纲**")
        elif current_stage_idx > 2:
            st.success("✔ **2. 制定大纲**")
        else:
            st.caption("2. 制定大纲")
    with col_s3:
        if current_stage_idx == 3:
            st.info("🔹 **3. 锁定确认**")
        elif current_stage_idx > 3:
            st.success("✔ **3. 锁定确认**")
        else:
            st.caption("3. 锁定确认")
    with col_s4:
        if current_stage_idx == 4:
            st.success("🎉 **4. 交付成稿**")
        else:
            st.caption("4. 交付成稿")
else:
    st.markdown(
        f"""
        <div style="background: {banner_bg}; border: 1px solid {banner_border}; border-radius: 12px; padding: 10px 16px; margin: 4px 0 14px 0; display: flex; align-items: center; justify-content: space-between; box-shadow: {card_shadow};">
            <div style="font-size: 13px; color: {text_main};">
                🌟 <b>自由灵感模式</b> · 由 <b>{current_expert.name}</b> 实时协同。支持发散创作与交互追问胶囊，随时一键直出交付物。
            </div>
            <div style="font-size: 12px; color: {text_sub};">
                大纲: {'✅ 已就绪' if state.draft else '⏳ 规划中'} ｜ 正文: {'🎉 已产出' if state.content else '⏳ 待撰写'}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

# ---------------------------------------------------------------------------
# 动态分栏布局：支持助手区与写作工作台切换 (2栏收拢 vs 3栏全开)
# ---------------------------------------------------------------------------
if st.session_state.show_workbench:
    chat_col, panel_col = st.columns([1, 1], gap="medium")
else:
    chat_col, panel_col = st.columns([1, 0.001], gap="small")


# ---------------------------------------------------------------------------
# 自由探索模式执行回调
# ---------------------------------------------------------------------------
def run_autonomous_turn(user_msg: str):
    user_display = user_msg.strip()
    turn_target = detect_target_kind(user_display, default=conversation_target)
    with chat_col:
        with st.chat_message("user", avatar="👤"):
            st.markdown(user_display)
        with st.chat_message("assistant", avatar=current_expert.avatar):
            with st.status("正在处理你的请求…", expanded=True) as turn_status:
                att_snippet = ""
                if st.session_state.get("current_attachment"):
                    att = st.session_state.current_attachment
                    att_snippet = att.get("text", att.get("text_snippet", ""))

                try:
                    res = st.session_state.autonomous_agent.step(
                        user_display,
                        target_kind=turn_target,
                        skill_id=skill_id.strip() or None,
                        history=st.session_state.messages,
                        current_outline=st.session_state.chat_state.draft,
                        current_content=st.session_state.chat_state.content,
                        attachment_text=att_snippet,
                        memory={**st.session_state.chat_state.autonomous_memory,
                                "requirements": st.session_state.chat_state.requirements},
                        progress=lambda label: turn_status.update(label=label),
                    )

                    st.session_state.chat_state.autonomous_memory = res.memory
                    st.session_state.chat_state.requirements = res.memory.get("requirements", {})
                    # 状态同步到界面看板
                    if res.outline:
                        res.outline["target_kind"] = res.target_kind
                        st.session_state.chat_state.draft = res.outline
                    if res.content:
                        res.content["target_kind"] = res.target_kind
                        st.session_state.chat_state.content = res.content
                    if res.rendered_file:
                        # Keep prior versions so a revision can be compared or downloaded.
                        st.session_state.downloads.append(res.rendered_file)

                    # 同步激活的选项式追问 (Option Chips)
                    st.session_state.active_follow_ups = res.follow_up_questions or []

                    # 格式化展示内容
                    asst_parts = []
                    asst_parts.append(res.reply_text)
                    if res.rendered_file:
                        kind_label = (
                            "Word 文档 (.docx)"
                            if res.target_kind == "document"
                            else ("演示文稿 (.pptx)" if res.target_kind == "presentation" else "文本文件")
                        )
                        asst_parts.append(
                            f"📦 **已完成文件渲染**：`{res.rendered_file['name']}`（{kind_label}，前往右侧看板「💾 文件下载」查收）"
                        )

                    if res.outline and not res.content:
                        asst_parts.append("\n".join(f"{i}. {n['title']}" for i, n in enumerate(res.outline["nodes"], 1)))
                    asst_msg = "\n\n".join(asst_parts)
                    turn_status.update(label="本轮完成", state="complete", expanded=False)
                    st.session_state.messages.append({"role": "user", "text": user_display})
                    st.session_state.messages.append(
                        {
                            "role": "assistant",
                            "text": asst_msg,
                            "avatar": current_expert.avatar,
                            "expert": current_expert.name,
                            "context_text": res.reply_text,
                            "follow_up_questions": res.follow_up_questions,
                        }
                    )
                    rerun_workbench()
                except Exception as exc:
                    turn_status.update(label="本轮未完成，原稿已保留", state="error")
                    err_msg = f"❌ 本轮处理失败: {exc}"
                    st.error(err_msg)
                    st.session_state.messages.append({"role": "user", "text": user_display})
                    st.session_state.messages.append({"role": "assistant", "text": err_msg, "avatar": "⚠️", "failed": True})
                    rerun_workbench()


# ---------------------------------------------------------------------------
# 专业模式动作执行回调
# ---------------------------------------------------------------------------
def run_action(
    operation: str,
    message: str = "",
    extra_payload: dict | None = None,
    override_target: str | None = None,
):
    extra = extra_payload or {}
    op_labels = {
        "clarify": "需求澄清",
        "create_outline": "生成大纲",
        "revise_outline": "修订大纲",
        "explain_outline": "解释大纲",
        "confirm_outline": "确认大纲",
        "generate": "生成内容",
    }
    user_display = message if message else f"👉 执行操作：【{op_labels.get(operation, operation)}】"

    # 获取当前挂载的原材料文件 ID（若有）
    current_file_ids = [st.session_state.current_attachment["file_id"]] if st.session_state.get("current_attachment") else []

    # 确定执行目标形态
    existing_target = (st.session_state.chat_state.draft or {}).get("target_kind")
    effective_target = (
        existing_target
        if (existing_target and operation != "create_outline")
        else (override_target or existing_target or target)
    )

    # 1. 在界面渲染用户发出的消息
    with chat_col:
        with st.chat_message("user", avatar="👤"):
            st.markdown(user_display)
        with st.chat_message("assistant", avatar=current_expert.avatar):
            with st.spinner(f"{current_expert.avatar} {current_expert.name} 正在推进【{op_labels.get(operation, operation)}】..."):
                try:
                    # Free-mode outlines have no backend draft/version identity.
                    # Bring their requirements into the guided workflow without
                    # pretending that a local outline was already confirmed there.
                    if state.draft and not state.draft.get("draft_id"):
                        saved_draft, saved_content = deepcopy(state.draft), deepcopy(state.content)
                        handoff = json.dumps({
                            "已确认的创作要求": state.requirements,
                            "当前大纲": saved_draft,
                            "用户本轮指令": message,
                        }, ensure_ascii=False)
                        try:
                            client.execute("clarify", target_kind=effective_target,
                                           message="从自由创作转入专业流程，保留下列要求与结构：" + handoff,
                                           state=state, current_file_ids=current_file_ids,
                                           skill_id=skill_id.strip() or None)
                        finally:
                            # Preserve the visible draft even if the remote call fails.
                            state.draft, state.content = saved_draft, saved_content
                        operation = "create_outline"
                        message = "请将现有结构整理为专业流程大纲，保留已确认要求：" + handoff
                    result = client.execute(
                        operation,
                        target_kind=effective_target,
                        message=message,
                        state=st.session_state.chat_state,
                        current_file_ids=current_file_ids,
                        draft_id=(st.session_state.chat_state.draft or {}).get("draft_id"),
                        expected_outline_version=(st.session_state.chat_state.draft or {}).get("outline_version"),
                        skill_id=skill_id.strip() or None,
                        **extra,
                    )
                    # 确保 draft 记录有效的 target_kind
                    if st.session_state.chat_state.draft and "target_kind" not in st.session_state.chat_state.draft:
                        st.session_state.chat_state.draft["target_kind"] = effective_target

                    # 同步附件下载缓存
                    for artifact in st.session_state.chat_state.artifacts:
                        fid = artifact.get("file_id")
                        fname = artifact.get("filename", fid)
                        if fid and not any(d["id"] == fid for d in st.session_state.downloads):
                            data = client.artifact(fid)
                            st.session_state.downloads.append({"id": fid, "name": fname, "data": data})

                    asst_text = format_assistant_message(st.session_state.chat_state, operation)
                    st.session_state.messages.append({"role": "user", "text": user_display})
                    st.session_state.messages.append(
                        {
                            "role": "assistant",
                            "text": asst_text,
                            "avatar": current_expert.avatar,
                            "expert": current_expert.name,
                        }
                    )
                    rerun_workbench()
                except ChatClientError as exc:
                    st.session_state.messages.append({"role": "user", "text": user_display})
                    att_name = (st.session_state.get("current_attachment") or {}).get("filename", "所选附件")

                    if exc.code == "ATTACHMENT_INVALID":
                        err_text = (
                            f"❌ **参考材料解析失败 (ATTACHMENT_INVALID)**\n\n"
                            f"后端无法解析所挂载的附件 `{att_name}`，主要原因包括：\n"
                            "1. **PDF 为图片/扫描件**：PDF 内部无文字图层（系统当前未开启 OCR 识别）；\n"
                            "2. **Word 文档无正文**：文档内无段落或表格文字；\n"
                            "3. **文件为空或字数超限**：文件为 0 字节，或提取文字超过上限。\n\n"
                            "👉 请移除该附件后直接输入纯文本继续。"
                        )
                        st.error(f"❌ 附件解析失败 ({att_name})")
                    elif exc.code == "MATERIAL_OUTLINE_CONFLICT":
                        outline_title = (st.session_state.chat_state.draft or {}).get("title", "当前大纲")
                        err_text = (
                            f"⚠️ **参考材料与大纲主题冲突**\n\n"
                            f"挂载的材料 `{att_name}` 与当前确认的大纲主题（「{outline_title}」）严重不符。"
                        )
                        st.error(f"⚠️ 材料与大纲冲突")
                    elif exc.code == "TOPIC_REQUIRED":
                        err_text = "⚠️ **尚未明确写作主题**：请在输入框中简述主题后再生成大纲。"
                        st.error("⚠️ 需要明确写作主题")
                    elif exc.code == "OUTLINE_NOT_CONFIRMED":
                        err_text = "⚠️ **大纲尚未确认锁定**：请先发送「确认大纲」锁定结构后，再生成内容。"
                        st.error("⚠️ 大纲尚未确认锁定")
                    elif exc.code == "OUTLINE_REQUIRED":
                        err_text = "⚠️ **大纲尚未生成**：请先执行生成大纲操作。"
                        st.error("⚠️ 大纲尚未生成")
                    else:
                        err_text = f"❌ 调用失败 ({exc.code or 'ERROR'}): {exc}"
                        st.error(f"❌ 请求失败: {exc}")

                    st.session_state.messages.append({"role": "assistant", "text": err_text, "avatar": "⚠️"})
                    rerun_workbench()


# ---------------------------------------------------------------------------
# 左侧：交互式对话区
# ---------------------------------------------------------------------------
with chat_col:
    st.subheader("💬 对话工作区")
    if is_autonomous and state.requirements:
        with st.expander("已记住的创作要求", expanded=False):
            requirement_labels = {"topic": "主题", "audience": "受众", "language": "语言", "length": "篇幅", "style": "风格", "purpose": "用途", "constraints": "其他要求"}
            for name, value in state.requirements.items():
                st.markdown(f"**{requirement_labels.get(name, name)}**：{value}")
            st.caption("可以直接在对话里修改或取消这些要求。")
    chat_container = st.container()

    with chat_container:
        if not st.session_state.messages:
            welcome_intro = (
                f"👋 您好！我是 **{current_expert.name}**（{current_expert.title}）。\n\n"
                f"{current_expert.intro}\n\n"
                f"💡 **快速开启灵感**：请在下方输入您的写作需求或主题（例如：*帮我写一份2026年微服务架构演进方案* 或 *做一份季度工作总结汇报PPT*），我将为您全力以赴！"
            )
            with st.chat_message("assistant", avatar=current_expert.avatar):
                st.markdown(welcome_intro)

        for item in st.session_state.messages:
            msg_avatar = item.get("avatar") or ("👤" if item["role"] == "user" else current_expert.avatar)
            with st.chat_message(item["role"], avatar=msg_avatar):
                st.markdown(item["text"])

        # 可选追问一次提交，避免选择一项就丢失其他问题。
        follow_ups = st.session_state.get("active_follow_ups", [])
        if is_autonomous and follow_ups:
            chips_title_color = "#0071e3" if not is_dark else "#2997ff"
            st.markdown(
                f"""
                <div style="background: rgba(0, 113, 227, 0.06); border: 1px solid rgba(0, 113, 227, 0.2); border-radius: 14px; padding: 12px 18px; margin: 16px 0 10px 0;">
                    <div style="font-size: 13px; font-weight: 600; color: {chips_title_color}; margin-bottom: 2px;">
                        💡 专家建议与快速选项 · 点击一键带入交互
                    </div>
                    <div style="font-size: 11px; color: {text_sub};">可直接点击备选标签快速回复，无需手动敲字：</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            answers = []
            with st.form(f"followups_{st.session_state.current_session_id}"):
                for q_idx, q_item in enumerate(follow_ups, 1):
                    q_text = q_item.get("question", "")
                    options = q_item.get("options", [])
                    question_key = hashlib.sha256(q_text.encode()).hexdigest()[:12]
                    if options:
                        answer = st.selectbox(q_text, ["暂不选择", *options],
                                              key=f"answer_{st.session_state.current_session_id}_{question_key}")
                        if answer != "暂不选择":
                            answers.append(f"针对【{q_text}】，我的选择是：{answer}")
                    else:
                        answer = st.text_input(q_text, key=f"answer_{st.session_state.current_session_id}_{question_key}")
                        if answer.strip():
                            answers.append(f"针对【{q_text}】，我的回答是：{answer.strip()}")
                submitted = st.form_submit_button("提交补充", use_container_width=True)
            if submitted and answers:
                run_autonomous_turn("\n".join(answers))
            elif submitted:
                st.caption("可以选择一项，也可以直接在聊天框继续交流。")

    # 快捷操作建议区（专业模式下渲染）
    guidance = state.guidance or {}
    questions = guidance.get("questions", [])
    proposals = guidance.get("proposals", [])

    if not is_autonomous and (questions or proposals):
        with st.expander("💡 快捷采纳建议", expanded=True):
            for q in questions:
                opts = q.get("options", [])
                if opts:
                    st.caption(f"针对问题「{q.get('text')}」，可直接选择推荐方案：")
                    opt_cols = st.columns(min(len(opts), 3))
                    for i, opt in enumerate(opts):
                        if opt_cols[i % len(opt_cols)].button(
                            opt, key=f"btn_opt_{q['question_id']}_{i}", use_container_width=True
                        ):
                            run_action("clarify", opt, {"answers": {q["question_id"]: opt}})

            for p in proposals:
                p_col1, p_col2 = st.columns([3, 1])
                p_col1.markdown(f"**建议设定** `{p.get('field')}` = `{p.get('value')}` *({p.get('reason')})*")
                if p_col2.button("✅ 采纳此建议", key=f"btn_prop_{p['proposal_id']}", use_container_width=True):
                    run_action(
                        "clarify",
                        f"采纳建议：{p.get('field')}为{p.get('value')}",
                        {"accepted_proposal_ids": [p["proposal_id"]]},
                    )

    st.write("")

    # -----------------------------------------------------------------------
    # 附件管理与上传区域
    # -----------------------------------------------------------------------
    if st.session_state.get("current_attachment"):
        att = st.session_state.current_attachment
        att_c1, att_c2 = st.columns([4, 1])
        att_c1.info(f"📎 **已挂载参考材料**: `{att['filename']}` ({att['size_str']})")
        if att_c2.button("🗑️ 移除附件", key="btn_remove_att", use_container_width=True, help="移除当前材料并恢复纯文本交互"):
            st.session_state.current_attachment = None
            st.session_state.uploader_key += 1
            rerun_workbench()

        # 智能快捷操作提示
        stem_name = Path(att["filename"]).stem
        target_name = "演示幻灯片 PPT" if target == "presentation" else ("Word 文档" if target == "document" else "纯正文")
        if is_autonomous:
            if st.button(
                f"⚡ 快捷指令：依据《{stem_name}》一键直接生成【{target_name}】",
                key="btn_quick_auto_gen_from_att",
                type="primary",
                use_container_width=True,
            ):
                run_autonomous_turn(f"请依据参考材料《{stem_name}》的核心内容，直接生成一份完整的【{target_name}】，并输出可编辑文件供下载。")
        else:
            if not state.draft:
                if st.button(
                    f"📑 快捷指令：以《{stem_name}》为核心主题制定大纲",
                    key="btn_quick_outline_from_att",
                    type="primary",
                    use_container_width=True,
                ):
                    run_action("create_outline", f"请以《{stem_name}》为主题，提炼核心结构制定大纲")

    with st.expander(
        "📎 上传参考原材料（TXT / MD / PDF / DOCX）",
        expanded=not bool(st.session_state.get("current_attachment")),
    ):
        st.caption("💡 提示：支持 .txt、.md、.docx、.pdf。PDF 需含可选中文字，单文件上限 20MB。")
        uploaded = st.file_uploader(
            "选择文件",
            type=["txt", "md", "pdf", "docx"],
            key=f"uploader_{st.session_state.current_session_id}_{st.session_state.uploader_key}",
            label_visibility="collapsed",
        )
        if uploaded is not None:
            file_bytes = uploaded.getvalue()
            file_hash = hashlib.sha256(uploaded.name.encode() + file_bytes).hexdigest()
            if (
                not st.session_state.get("current_attachment")
                or st.session_state.current_attachment.get("hash") != file_hash
            ):
                ok, reason = validate_attachment(uploaded.name, file_bytes)
                if not ok:
                    st.error(f"⚠️ 文件预检未通过: {reason}")
                else:
                    with st.spinner(f"正在上传材料「{uploaded.name}」到 BFF 网关..."):
                        try:
                            attachment_text = read_attachment_bytes(uploaded.name, file_bytes)
                            uploaded_res = client.upload(
                                uploaded.name, file_bytes, uploaded.type, session_id=state.session
                            )
                            size_kb = max(round(len(file_bytes) / 1024, 1), 0.1)

                            st.session_state.current_attachment = {
                                "file_id": uploaded_res["file_id"],
                                "filename": uploaded.name,
                                "size_str": f"{size_kb} KB",
                                "hash": file_hash,
                                "text": attachment_text,
                                "char_count": len(attachment_text),
                            }
                            st.success(f"✅ 已读取材料: {uploaded.name} · {len(attachment_text)} 字符（含可解析表格）")
                            rerun_workbench()
                        except Exception as upload_err:
                            st.error(f"❌ 材料上传或解析失败，未替换现有材料: {upload_err}")

    # 专业模式操作快捷指令栏
    if not is_autonomous:
        st.caption("手动快捷动作（亦可直接在下方对话框用自然语言表达，模型将自动判断并流转）：")
        operations = ["clarify", "create_outline", "revise_outline", "explain_outline", "confirm_outline", "generate"]
        operation_labels = {
            "clarify": "💬 补充澄清",
            "create_outline": "📑 生成大纲",
            "revise_outline": "✍️ 修订大纲",
            "explain_outline": "❓ 解释大纲",
            "confirm_outline": "🔒 确认大纲",
            "generate": "🚀 生成内容",
        }

        confirmed = guidance.get("stage") == "confirmed" or bool(state.draft and state.draft.get("confirmed_hash"))
        enabled = {
            "clarify": True,
            "create_outline": True,
            "revise_outline": bool(state.draft),
            "explain_outline": bool(state.draft),
            "confirm_outline": bool(state.draft) and not confirmed,
            "generate": bool(state.draft) and confirmed,
        }

        btn_cols = st.columns(6)
        chosen_op = None
        for op, col in zip(operations, btn_cols):
            is_primary = (
                (op == "create_outline" and current_stage_idx == 1)
                or (op == "confirm_outline" and current_stage_idx == 2)
                or (op == "generate" and current_stage_idx == 3)
            )
            if col.button(
                operation_labels[op],
                key=f"op_{op}",
                disabled=not enabled[op],
                type="primary" if is_primary else "secondary",
                use_container_width=True,
            ):
                chosen_op = op
    else:
        chosen_op = None

    # 聊天输入辅助栏（包含打开/收起写作工作台按钮，快捷附件与操作）
    col_in_btn1, col_in_btn2 = st.columns([1, 1])
    with col_in_btn1:
        wb_btn_label = "📖 收起写作工作台" if st.session_state.show_workbench else "✨ 打开写作工作台"
        wb_btn_help = "点击在右侧展开沉浸式写作工作台画布（整体自适应分成3栏）" if not st.session_state.show_workbench else "点击收起右侧工作台，扩大对话区空间"
        if st.button(wb_btn_label, key="btn_toggle_workbench", use_container_width=True, type="primary" if not st.session_state.show_workbench else "secondary", help=wb_btn_help):
            st.session_state.show_workbench = not st.session_state.show_workbench
            rerun_workbench()

    with col_in_btn2:
        target_name = "Word 文档 (.docx)" if target == "document" else ("幻灯片 (.pptx)" if target == "presentation" else "纯正文")
        st.caption(f"创作目标：**{target_name}** ｜ 顾问：**{current_expert.name}**")

    # 聊天输入框
    prompt = st.chat_input("输入你的需求、修改意见、自由探讨，或直接下达创作指令...")

    if is_autonomous:
        if prompt:
            run_autonomous_turn(prompt)
    else:
        if prompt or chosen_op:
            op_to_run = chosen_op
            msg_to_send = (prompt or "").strip()
            override_target = None

            # 针对用户在对话框直接输入的自然语言，使用大模型意图路由器智能决策状态流转！
            if not op_to_run and prompt:
                with st.spinner("🤖 AI 正在智能研判意图并驱动状态流转..."):
                    route_res = route_professional_intent(
                        prompt,
                        stage=(state.guidance or {}).get("stage", "clarify"),
                        has_draft=bool(state.draft),
                        is_confirmed=(state.guidance or {}).get("stage") == "confirmed"
                        or bool(state.draft and state.draft.get("confirmed_hash")),
                        current_title=(state.draft or {}).get("title", ""),
                        agent=st.session_state.autonomous_agent,
                    )
                    op_to_run = route_res.get("operation", "clarify")
                    msg_to_send = route_res.get("clean_message", prompt)
                    override_target = route_res.get("override_target")

                    op_name = {
                        "clarify": "需求澄清",
                        "create_outline": "生成大纲",
                        "revise_outline": "修订大纲",
                        "explain_outline": "解释大纲",
                        "confirm_outline": "确认大纲",
                        "generate": "生成内容",
                    }.get(op_to_run, op_to_run)

                    st.toast(f"🤖 AI 语义研判：自动流转至【{op_name}】")

            # 针对 create_outline 检查是否缺少主题与意图形态推断
            if op_to_run == "create_outline":
                if msg_to_send:
                    override_target = detect_target_kind(msg_to_send, default=target)
                current_topic = getattr(state, "requirements", {}).get("topic", "")
                if not current_topic and not msg_to_send:
                    att = st.session_state.get("current_attachment")
                    if att:
                        stem_name = Path(att["filename"]).stem
                        msg_to_send = f"请以参考材料《{stem_name}》为核心主题制定大纲"
                        override_target = detect_target_kind(msg_to_send, default=target)
                    else:
                        st.warning("⚠️ 制定大纲需要明确主题。请在对话框中输入您想要的主题，再继续生成大纲。")
                        st.stop()

            # 确认与生成阶段不需要额外用户指令
            if op_to_run in {"confirm_outline", "generate"}:
                msg_to_send = ""

            run_action(op_to_run, msg_to_send, override_target=override_target)


# ---------------------------------------------------------------------------
# 右侧：实时产物写作工作台画布看板 (仅在用户打开工作台时展开为第 3 栏)
# ---------------------------------------------------------------------------
if st.session_state.show_workbench:
    with panel_col:
        col_wb_head1, col_wb_head2 = st.columns([4, 1])
        with col_wb_head1:
            st.subheader("📑 写作工作台")
        with col_wb_head2:
            if st.button("✕ 收起", key="btn_close_panel", help="收起写作工作台"):
                st.session_state.show_workbench = False
                rerun_workbench()

        tab_outline, tab_content, tab_download, tab_council = st.tabs(
            ["📑 结构大纲", "📖 交付正文", "💾 文件下载", "👥 专家智囊团"]
        )

        with tab_outline:
            if state.draft:
                draft = state.draft
                is_locked = bool(draft.get("confirmed_hash"))
                st.markdown(f"### 📑 {draft.get('title', '未命名大纲')}")
                badge_confirmed = "🔒 **已锁定确认**" if is_locked else "⏳ **草稿待确认**"
                st.caption(
                    f"版本: `v{draft.get('outline_version', 1)}` | 目标形态: `{draft.get('target_kind', target)}` | 状态: {badge_confirmed}"
                )

                nodes = draft.get("nodes", [])
                if not nodes:
                    st.info("大纲中暂无章节节点")
                for idx, node in enumerate(nodes, 1):
                    nid = node.get("node_id", f"node-{idx}")
                    ntitle = node.get("title", f"第 {idx} 节")
                    st.markdown(
                        f"""
                        <div style="background: {card_bg}; border: 1px solid {card_border}; border-radius: 12px; padding: 12px 16px; margin-bottom: 8px; display: flex; justify-content: space-between; align-items: center; box-shadow: {card_shadow};">
                            <span style="font-weight: 600; font-size: 14px; color: {text_main};">{idx}. 📌 {ntitle}</span>
                            <code style="font-size: 11px; background: {tag_bg}; border: 1px solid {tag_border}; padding: 3px 8px; border-radius: 6px; color: {text_main}; font-weight: 500;">{nid}</code>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
            else:
                st.info("💡 尚未生成大纲。请在左侧对话区提出写作构想或上传参考材料。")

        with tab_content:
            if state.content:
                content = state.content
                c_title = content.get("title", "未命名文章")
                c_summary = content.get("summary", "")
                sections = content.get("sections", [])
                sec_count = len(sections)

                # 拼装纯文本正文用于快捷一键复制
                full_text_lines = [f"# {c_title}\n"]
                if c_summary:
                    full_text_lines.append(f"> 摘要：{c_summary}\n")
                for idx, sec in enumerate(sections, 1):
                    full_text_lines.append(f"## {idx}. {sec.get('title', '')}\n")
                    if sec.get("body"):
                        full_text_lines.append(f"{sec.get('body')}\n")
                    if sec.get("notes"):
                        full_text_lines.append(f"> 备注说明：{sec.get('notes')}\n")
                full_article_text = "\n".join(full_text_lines)

                # 画布工具栏
                col_bar1, col_bar2 = st.columns([3, 1])
                with col_bar1:
                    st.caption(f"📊 正文统计：约 {len(full_article_text)} 字 ｜ 共 {sec_count} 个章节")
                with col_bar2:
                    with st.popover("📋 复制全文", help="点击弹出全文快速复制框"):
                        st.text_area("直接全选复制 (Command/Ctrl + A)", value=full_article_text, height=240)

                # 沉浸式豆包纸质白板画布排版
                st.markdown(
                    f"""
                    <div class="doubao-paper-canvas">
                        <div class="doubao-doc-header">
                            <div class="doubao-doc-title">📖 {c_title}</div>
                            <div class="doubao-doc-meta">
                                <span>字数: <b>{len(full_article_text)}</b></span>
                                <span>•</span>
                                <span>章节: <b>{sec_count}</b> 节</span>
                                <span>•</span>
                                <span>形态: <b>{target}</b></span>
                            </div>
                        </div>
                    """,
                    unsafe_allow_html=True,
                )

                if c_summary:
                    st.info(f"**💡 核心摘要与主旨**：\n\n{c_summary}")

                for idx, sec in enumerate(sections, 1):
                    sec_title = sec.get("title", f"第 {idx} 节")
                    sec_body = sec.get("body", "")
                    st.markdown(
                        f"""
                        <div class="doubao-section-block">
                            <div class="doubao-section-title">{idx}. 📌 {sec_title}</div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                    if sec_body:
                        st.markdown(sec_body)
                    notes = sec.get("notes", "")
                    if notes:
                        with st.expander(f"📝 查看第 {idx} 节演说备注与要点解析", expanded=False):
                            st.caption(notes)
                    st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)

                st.markdown("</div>", unsafe_allow_html=True)
            else:
                st.info("💡 正文内容尚未生成。在左侧对话中与专家确定大纲后，AI 将自动出稿并排版呈现。")

        with tab_download:
            st.markdown("### 💾 交付成果文件导出")
            if not st.session_state.downloads:
                st.info("暂无生成好的文件可供下载。")
            else:
                for item in st.session_state.downloads:
                    fname = item["name"]
                    fdata = item["data"]
                    size_kb = max(round(len(fdata) / 1024, 1), 0.1)
                    is_docx = fname.endswith(".docx")
                    is_pptx = fname.endswith(".pptx")
                    icon = "📄" if is_docx else ("📊" if is_pptx else "📝")
                    badge_type = "Word 文档 (.docx)" if is_docx else ("演示幻灯片 (.pptx)" if is_pptx else "文本文件")

                    st.markdown(
                        f"""
                        <div style="background: {card_bg}; border: 1px solid {card_border}; border-radius: 14px; padding: 14px 18px; margin-bottom: 12px; display: flex; align-items: center; justify-content: space-between; box-shadow: {card_shadow};">
                            <div>
                                <div style="font-size: 15px; font-weight: 600; color: {text_main};">{icon} {fname}</div>
                                <div style="font-size: 12px; color: {text_sub}; margin-top: 3px;">格式: {badge_type} ｜ 大小: {size_kb} KB</div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                    mime = (
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        if is_docx
                        else (
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                            if is_pptx
                            else "application/octet-stream"
                        )
                    )
                    st.download_button(
                        label=f"⬇️ 立即下载 {fname}",
                        data=fdata,
                        file_name=fname,
                        mime=mime,
                        key=f"dl_btn_{st.session_state.current_session_id}_{item['id']}",
                        use_container_width=True,
                        type="primary",
                    )

        with tab_council:
            st.markdown("### 👥 专家顾问团队成员名录")
            st.caption("GenSlide AI 写作工作台联合业界专家顾问，提供各细分领域的专业能力支撑：")

            for exp in all_experts:
                is_current = exp.id == current_expert.id
                border_style = (
                    "border: 1.5px solid #0071e3; box-shadow: 0 4px 18px rgba(0, 113, 227, 0.15);"
                    if is_current
                    else f"border: 1px solid {card_border};"
                )
                current_tag = '<span style="background: #0071e3; color: white; padding: 2px 8px; border-radius: 9999px; font-size: 11px; font-weight: 600;">🟢 当前坐镇</span>' if is_current else ""
                t_html = " ".join(
                    f'<span style="background: {tag_bg}; color: {text_main}; border: 1px solid {tag_border}; border-radius: 4px; padding: 2px 8px; font-size: 10px; font-weight: 500;">{t}</span>'
                    for t in exp.tags
                )

                st.markdown(
                    f"""
                    <div style="background: {card_bg}; {border_style} border-radius: 14px; padding: 14px 16px; margin-bottom: 12px; box-shadow: {card_shadow};">
                        <div style="display: flex; align-items: center; justify-content: space-between;">
                            <div style="display: flex; align-items: center; gap: 10px;">
                                <span style="font-size: 24px;">{exp.avatar}</span>
                                <div>
                                    <strong style="font-size: 14px; color: {text_main};">{exp.name}</strong>
                                    <div style="font-size: 11px; color: {text_sub};">{exp.title}</div>
                                </div>
                            </div>
                            <div>{current_tag}</div>
                        </div>
                        <div style="font-size: 12px; color: {text_caption}; margin-top: 8px; line-height: 1.5;">
                            {exp.intro}
                        </div>
                        <div style="margin-top: 8px; display: flex; flex-wrap: wrap; gap: 4px;">
                            {t_html}
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )


save_session(st.session_state)
